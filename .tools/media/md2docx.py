# -*- coding: utf-8 -*-
"""Markdown → DOCX 转换（面向本项目文档：标题/段落/表格/列表/引用/代码块/图片/行内强调）。
用法：python md2docx.py <input.md> [output.docx]
"""
import io, os, re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from PIL import Image

MAX_IMG_W_CM = 15.0

def set_cjk(style_or_run, latin='Calibri', cjk='微软雅黑', size=None, bold=None, color=None):
    font = style_or_run.font
    font.name = latin
    if size: font.size = Pt(size)
    if bold is not None: font.bold = bold
    if color: font.color.rgb = RGBColor(*color)
    rpr = style_or_run._element.get_or_add_rPr() if hasattr(style_or_run, '_element') else None
    if rpr is not None:
        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is None:
            rfonts = rpr.makeelement(qn('w:rFonts'), {}); rpr.append(rfonts)
        rfonts.set(qn('w:eastAsia'), cjk)

def add_runs(par, text, base_size=10.5):
    """处理 **bold**、`code`、*italic* 行内标记。"""
    tokens = re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)', text)
    for tok in tokens:
        if not tok: continue
        if tok.startswith('**') and tok.endswith('**'):
            r = par.add_run(tok[2:-2]); r.bold = True; set_cjk(r, size=base_size)
        elif tok.startswith('`') and tok.endswith('`'):
            r = par.add_run(tok[1:-1]); r.font.name = 'Consolas'; set_cjk(r, latin='Consolas', cjk='微软雅黑', size=base_size - 0.5)
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.italic = True; set_cjk(r, size=base_size)
        else:
            r = par.add_run(tok); set_cjk(r, size=base_size)

def convert(md_path, out_path):
    src = io.open(md_path, encoding='utf-8').read()
    base = os.path.dirname(os.path.abspath(md_path))
    doc = Document()
    # 页边距
    for s in doc.sections:
        s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.2); s.left_margin = Cm(2.6); s.right_margin = Cm(2.6)
    normal = doc.styles['Normal']; set_cjk(normal, size=10.5)
    lines = src.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        # 代码块
        if line.startswith('```'):
            i += 1; buf = []
            while i < len(lines) and not lines[i].startswith('```'):
                buf.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            r = p.add_run('\n'.join(buf)); r.font.name = 'Consolas'; r.font.size = Pt(9)
            set_cjk(r, latin='Consolas', size=9)
            continue
        # 图片
        m = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', line)
        if m:
            img_path = os.path.normpath(os.path.join(base, m.group(2)))
            if os.path.exists(img_path):
                with Image.open(img_path) as im:
                    w, h = im.size
                w_cm = min(MAX_IMG_W_CM, w / 96.0 * 2.54)
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(img_path, width=Cm(w_cm))
                if m.group(1):
                    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = cap.add_run(m.group(1)); r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
                    set_cjk(r, size=9)
            i += 1; continue
        # 表格
        if line.startswith('|') and i + 1 < len(lines) and set(lines[i + 1].replace(' ', '')) <= set('|-:'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                rows.append(cells); i += 1
            rows = [rows[0]] + rows[2:]
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=0, cols=ncol); t.style = 'Table Grid'
            for ri, row in enumerate(rows):
                cells = t.add_row().cells
                for ci in range(ncol):
                    txt = row[ci] if ci < len(row) else ''
                    par = cells[ci].paragraphs[0]
                    add_runs(par, txt, base_size=9.5)
                    if ri == 0:
                        for r in par.runs: r.bold = True
            doc.add_paragraph()
            continue
        # 标题
        m = re.match(r'^(#{1,4})\s+(.*)$', line)
        if m:
            lvl = len(m.group(1)); text = m.group(2)
            h = doc.add_heading('', level=min(lvl, 4))
            add_runs(h, text, base_size={1: 18, 2: 15, 3: 13, 4: 11.5}[min(lvl, 4)])
            for r in h.runs:
                r.font.color.rgb = RGBColor(0x1F, 0x35, 0x64)
            i += 1; continue
        # 引用
        if line.startswith('>'):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            r = p.add_run(line.lstrip('> ').strip()); r.italic = True; r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            set_cjk(r, size=10)
            i += 1; continue
        # 分隔线
        if re.match(r'^-{3,}$', line):
            doc.add_paragraph('─' * 40).alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue
        # 列表
        m = re.match(r'^\s*[-*•]\s+(.*)$', line)
        if m:
            p = doc.add_paragraph(style='List Bullet'); add_runs(p, m.group(1)); i += 1; continue
        m = re.match(r'^\s*\d+[.)]\s+(.*)$', line)
        if m:
            p = doc.add_paragraph(style='List Number'); add_runs(p, m.group(1)); i += 1; continue
        # 空行
        if not line.strip():
            i += 1; continue
        p = doc.add_paragraph(); add_runs(p, line)
        i += 1
    doc.save(out_path)
    print('written', out_path)

if __name__ == '__main__':
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(src)[0] + '.docx'
    convert(src, dst)
